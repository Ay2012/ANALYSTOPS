"""Deterministic Bronze intake validation for submitted workbooks."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_RESULTS_DIR = PROJECT_ROOT / "data" / "ingestion" / "results"

EXPECTED_COLUMNS = (
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
)
REQUIRED_COLUMNS = (
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Country",
)
EXPECTED_FIELD = {
    "Invoice": "invoice",
    "StockCode": "product",
    "Description": "description",
    "Quantity": "quantity",
    "InvoiceDate": "date",
    "Price": "price",
    "Customer ID": "customer",
    "Country": "country",
}
FIELD_ALIASES = {
    "invoice": {"invoice", "invoiceid", "invoiceno", "order", "orderid"},
    "product": {"stockcode", "sku", "product", "productid", "item", "itemid"},
    "description": {"description", "productdescription", "itemdescription"},
    "quantity": {"quantity", "qty", "units", "unitssold"},
    "date": {"invoicedate", "transactiondate", "orderdate", "date"},
    "price": {"price", "unitprice", "amount", "salesamount"},
    "customer": {"customerid", "customernumber", "customer"},
    "country": {"country", "market", "region"},
}
TRANSACTION_ROLES = {
    "transaction_identifier": {"invoice"},
    "product_identifier": {"product"},
    "activity_measure": {"quantity", "price"},
    "event_time": {"date"},
}
QUALITY_RANK = {"PASS": 0, "WARN": 1, "REVIEW": 2, "BLOCK": 3}
DUPLICATE_CODES = {"duplicate_file_hash", "duplicate_content_fingerprint"}
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_SHEETS = 50
MAX_ROWS_PER_SHEET = 1_000_000
MAX_COLUMNS_PER_SHEET = 1_000
MAX_REPORTED_LOCATIONS = 20
DUPLICATE_REVIEW_RATE = 0.10


@dataclass(frozen=True)
class IntakeResult:
    file_path: str
    file_hash: str | None
    content_fingerprint: str | None
    lifecycle_state: str
    quality_disposition: str
    decision: str
    selected_sheet: str | None
    sheet_candidates: list[dict[str, object]]
    observed_schema: list[str]
    row_count: int
    reason_codes: list[str]
    findings: list[dict[str, object]]

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


def validate_workbook(
    path: Path | str,
    *,
    seen_hashes: Iterable[str] = (),
    seen_content_fingerprints: Iterable[str] = (),
    baseline_row_count: int | None = None,
    selected_sheet: str | None = None,
) -> IntakeResult:
    """Return an intake decision and evidence for one workbook."""

    workbook_path = Path(path)
    file_hash = _file_hash(workbook_path) if workbook_path.is_file() else None
    seen_hashes = set(seen_hashes)
    seen_content_fingerprints = set(seen_content_fingerprints)
    findings: list[dict[str, object]] = []

    if file_hash and file_hash in seen_hashes:
        findings.append(_finding("duplicate_file_hash", "PASS"))

    findings.extend(_package_safety_findings(workbook_path))
    if _has_blocking_finding(findings):
        return _result(
            file_path=str(workbook_path),
            file_hash=file_hash,
            content_fingerprint=None,
            selected_sheet=None,
            sheet_candidates=[],
            observed_schema=[],
            row_count=0,
            findings=findings,
        )

    try:
        findings.extend(_workbook_safety_findings(workbook_path))
    except Exception as exc:
        return _result(
            file_path=str(workbook_path),
            file_hash=file_hash,
            content_fingerprint=None,
            selected_sheet=None,
            sheet_candidates=[],
            observed_schema=[],
            row_count=0,
            findings=findings
            + [_finding("workbook_read_failed", "BLOCK", error=str(exc))],
        )

    if _has_blocking_finding(findings):
        return _result(
            file_path=str(workbook_path),
            file_hash=file_hash,
            content_fingerprint=None,
            selected_sheet=None,
            sheet_candidates=[],
            observed_schema=[],
            row_count=0,
            findings=findings,
        )

    try:
        sheets = pd.read_excel(workbook_path, sheet_name=None)
    except Exception as exc:
        return _result(
            file_path=str(workbook_path),
            file_hash=file_hash,
            content_fingerprint=None,
            selected_sheet=None,
            sheet_candidates=[],
            observed_schema=[],
            row_count=0,
            findings=findings
            + [_finding("workbook_read_failed", "BLOCK", error=str(exc))],
        )

    candidates = [_sheet_candidate(name, frame) for name, frame in sheets.items()]
    plausible = [item for item in candidates if item["plausible"]]
    if not plausible:
        return _result(
            file_path=str(workbook_path),
            file_hash=file_hash,
            content_fingerprint=None,
            selected_sheet=None,
            sheet_candidates=candidates,
            observed_schema=[],
            row_count=0,
            findings=findings + [_finding("no_plausible_transaction_sheet", "BLOCK")],
        )

    plausible.sort(key=lambda item: (-int(item["score"]), str(item["sheet"])))
    if selected_sheet is None:
        selected = plausible[0]
    else:
        selected = next(
            (item for item in plausible if item["sheet"] == selected_sheet),
            None,
        )
        if selected is None:
            return _result(
                file_path=str(workbook_path),
                file_hash=file_hash,
                content_fingerprint=None,
                selected_sheet=None,
                sheet_candidates=candidates,
                observed_schema=[],
                row_count=0,
                findings=findings
                + [
                    _finding(
                        "selected_sheet_not_plausible",
                        "BLOCK",
                        selected_sheet=selected_sheet,
                    )
                ],
            )
    chosen_sheet = str(selected["sheet"])
    frame = sheets[chosen_sheet]
    content_fingerprint = _content_fingerprint(frame)

    if len(sheets) > 1:
        findings.append(
            _finding(
                "extra_sheets_present",
                "WARN",
                extra_sheets=[name for name in sheets if name != chosen_sheet],
            )
        )
    if (
        selected_sheet is None
        and len(plausible) > 1
        and plausible[1]["score"] == selected["score"]
    ):
        findings.append(_finding("ambiguous_transaction_sheets", "REVIEW"))
    if content_fingerprint in seen_content_fingerprints:
        findings.append(_finding("duplicate_content_fingerprint", "PASS"))

    findings.extend(_schema_findings(frame))
    findings.extend(_type_findings(frame))
    findings.extend(_duplicate_findings(frame))
    findings.extend(
        _row_count_findings(
            frame,
            baseline_row_count,
            workbook_name=workbook_path.name,
            sheet_name=chosen_sheet,
        )
    )
    findings.extend(_safety_findings(frame, chosen_sheet))

    return _result(
        file_path=str(workbook_path),
        file_hash=file_hash,
        content_fingerprint=content_fingerprint,
        selected_sheet=chosen_sheet,
        sheet_candidates=candidates,
        observed_schema=[str(column) for column in frame.columns],
        row_count=len(frame),
        findings=findings,
    )


def write_result(result: IntakeResult, output_dir: Path | str = DEFAULT_RESULTS_DIR) -> Path:
    output_path = Path(output_dir) / f"{Path(result.file_path).stem}.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result.to_dict(), indent=2, sort_keys=True) + "\n")
    return output_path


def _sheet_candidate(sheet: str, frame: pd.DataFrame) -> dict[str, object]:
    matched_field_set = _matched_fields(frame.columns)
    matched_fields = sorted(matched_field_set)
    matched_roles = sorted(
        role
        for role, fields in TRANSACTION_ROLES.items()
        if matched_field_set.intersection(fields)
    )
    missing_roles = sorted(set(TRANSACTION_ROLES) - set(matched_roles))
    score = len(matched_fields)
    if len(frame) > 0:
        score += 1
    if len(frame.columns) >= 4:
        score += 1

    return {
        "sheet": sheet,
        "row_count": len(frame),
        "column_count": len(frame.columns),
        "score": score,
        "matched_fields": matched_fields,
        "matched_roles": matched_roles,
        "missing_roles": missing_roles,
        "plausible": not missing_roles,
    }


def _schema_findings(frame: pd.DataFrame) -> list[dict[str, object]]:
    findings: list[dict[str, object]] = []
    columns = {str(column) for column in frame.columns}
    matched_fields = _matched_fields(frame.columns)
    missing = [column for column in REQUIRED_COLUMNS if column not in columns]
    genuinely_missing = [
        column for column in missing if EXPECTED_FIELD[column] not in matched_fields
    ]
    renamed = [column for column in missing if column not in genuinely_missing]
    unexpected = [column for column in frame.columns if str(column) not in EXPECTED_COLUMNS]

    if genuinely_missing:
        findings.append(
            _finding("missing_required_columns", "BLOCK", columns=genuinely_missing)
        )
    if renamed:
        findings.append(_finding("renamed_required_columns", "REVIEW", columns=renamed))
    if unexpected:
        findings.append(
            _finding(
                "unexpected_columns",
                "WARN",
                columns=[str(column) for column in unexpected],
            )
        )

    return findings


def _type_findings(frame: pd.DataFrame) -> list[dict[str, object]]:
    findings: list[dict[str, object]] = []
    checks = {
        "Quantity": ("quantity_parse_failures", pd.to_numeric),
        "Price": ("price_parse_failures", pd.to_numeric),
        "InvoiceDate": ("date_parse_failures", pd.to_datetime),
    }

    for column, (code, parser) in checks.items():
        if column not in frame:
            continue
        parsed = parser(frame[column], errors="coerce")
        failures = int((frame[column].notna() & parsed.isna()).sum())
        if failures:
            findings.append(_finding(code, "REVIEW", column=column, count=failures))

    if "InvoiceDate" in frame:
        slash_dates = int(
            frame["InvoiceDate"]
            .dropna()
            .map(lambda value: isinstance(value, str) and "/" in value)
            .sum()
        )
        if slash_dates:
            findings.append(
                _finding("non_iso_date_strings", "REVIEW", column="InvoiceDate", count=slash_dates)
            )

    return findings


def _duplicate_findings(frame: pd.DataFrame) -> list[dict[str, object]]:
    duplicate_rows = int(frame.duplicated().sum())
    if not duplicate_rows:
        return []
    duplicate_rate = duplicate_rows / len(frame)
    disposition = "REVIEW" if duplicate_rate > DUPLICATE_REVIEW_RATE else "WARN"
    return [
        _finding(
            "exact_duplicate_rows",
            disposition,
            count=duplicate_rows,
            rate=round(duplicate_rate, 6),
            review_threshold=DUPLICATE_REVIEW_RATE,
        )
    ]


def _row_count_findings(
    frame: pd.DataFrame,
    baseline_row_count: int | None,
    *,
    workbook_name: str,
    sheet_name: str,
) -> list[dict[str, object]]:
    if len(frame) == 0:
        disposition = (
            "REVIEW" if baseline_row_count and baseline_row_count > 0 else "WARN"
        )
        return [
            _finding(
                "empty_transaction_sheet",
                disposition,
                workbook_name=workbook_name,
                sheet_name=sheet_name,
                row_count=0,
                message=(
                    f"{workbook_name} contains a valid transaction sheet "
                    "but no transaction records."
                ),
            )
        ]
    if baseline_row_count and len(frame) <= baseline_row_count * 0.5:
        return [
            _finding(
                "unexpectedly_low_row_count",
                "REVIEW",
                row_count=len(frame),
                baseline_row_count=baseline_row_count,
            )
        ]
    return []


def _safety_findings(
    frame: pd.DataFrame, sheet_name: str
) -> list[dict[str, object]]:
    formula_locations: list[str] = []
    prompt_locations: list[str] = []
    prompt_pattern = re.compile(
        r"ignore (?:all )?(?:previous|prior) instructions|system prompt|"
        r"\bact as (?:an? |the )?(?:system|assistant|developer)\b|"
        r"\bcall (?:a |the )?tool\b|"
        r"\bexecute (?:this|the) (?:command|code|script|tool)\b|"
        r"\breveal (?:credentials|secrets|the prompt)\b",
        re.IGNORECASE,
    )

    for column_index in range(len(frame.columns)):
        series = frame.iloc[:, column_index]
        if series.dtype != object:
            continue
        for row_number, value in enumerate(series, start=2):
            if not isinstance(value, str):
                continue
            candidate = value.lstrip("\t\r")
            formula_like = candidate.startswith(("=", "+", "@")) or (
                candidate.startswith("-")
                and pd.isna(pd.to_numeric(candidate, errors="coerce"))
            )
            location = (
                f"{sheet_name}!{get_column_letter(column_index + 1)}{row_number}"
            )
            if formula_like or value.startswith(("\t", "\r")):
                formula_locations.append(location)
            if prompt_pattern.search(value):
                prompt_locations.append(location)

    findings = []
    if formula_locations:
        findings.append(
            _finding(
                "formula_like_text_cells",
                "WARN",
                count=len(formula_locations),
                locations=formula_locations[:MAX_REPORTED_LOCATIONS],
            )
        )
    if prompt_locations:
        findings.append(
            _finding(
                "prompt_injection_text",
                "WARN",
                count=len(prompt_locations),
                locations=prompt_locations[:MAX_REPORTED_LOCATIONS],
            )
        )
    return findings


def _package_safety_findings(path: Path) -> list[dict[str, object]]:
    if not path.exists():
        return []
    if path.stat().st_size > MAX_FILE_BYTES:
        return [
            _finding(
                "file_size_limit_exceeded",
                "BLOCK",
                bytes=path.stat().st_size,
                limit=MAX_FILE_BYTES,
            )
        ]
    if not zipfile.is_zipfile(path):
        return []

    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
    names = [entry.filename.lower() for entry in entries]
    total_compressed = sum(entry.compress_size for entry in entries)
    total_uncompressed = sum(entry.file_size for entry in entries)
    findings: list[dict[str, object]] = []

    if any(
        name.startswith("/")
        or ".." in PurePosixPath(name.replace("\\", "/")).parts
        for name in names
    ):
        findings.append(_finding("unsafe_archive_path", "BLOCK"))
    if total_uncompressed > MAX_UNCOMPRESSED_BYTES or (
        total_uncompressed / max(total_compressed, 1) > MAX_COMPRESSION_RATIO
    ):
        findings.append(
            _finding(
                "suspicious_archive_compression",
                "BLOCK",
                compressed_bytes=total_compressed,
                uncompressed_bytes=total_uncompressed,
            )
        )
    if any(name.endswith("vbaproject.bin") or "/macrosheets/" in name for name in names):
        findings.append(_finding("macro_content_present", "BLOCK"))
    if any("/activex/" in name or "/embeddings/" in name for name in names):
        findings.append(_finding("embedded_active_content_present", "BLOCK"))
    if any(
        "/externallinks/" in name or name.endswith("/connections.xml")
        for name in names
    ):
        findings.append(_finding("external_connections_present", "BLOCK"))
    return findings


def _workbook_safety_findings(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            return [
                _finding(
                    "sheet_count_limit_exceeded",
                    "BLOCK",
                    count=len(workbook.sheetnames),
                    limit=MAX_SHEETS,
                )
            ]

        formula_locations: list[str] = []
        external_formula_locations: list[str] = []
        findings: list[dict[str, object]] = []
        for sheet in workbook.worksheets:
            if (
                sheet.max_row > MAX_ROWS_PER_SHEET
                or sheet.max_column > MAX_COLUMNS_PER_SHEET
            ):
                findings.append(
                    _finding(
                        "worksheet_dimension_limit_exceeded",
                        "BLOCK",
                        sheet=sheet.title,
                        rows=sheet.max_row,
                        columns=sheet.max_column,
                    )
                )
                continue

            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    location = f"{sheet.title}!{cell.coordinate}"
                    if _is_external_formula(str(cell.value or "")):
                        external_formula_locations.append(location)
                    else:
                        formula_locations.append(location)

        if external_formula_locations:
            findings.append(
                _finding(
                    "external_formula_cells",
                    "BLOCK",
                    count=len(external_formula_locations),
                    locations=external_formula_locations[:MAX_REPORTED_LOCATIONS],
                )
            )
        if formula_locations:
            findings.append(
                _finding(
                    "formula_cells",
                    "REVIEW",
                    count=len(formula_locations),
                    locations=formula_locations[:MAX_REPORTED_LOCATIONS],
                )
            )
        return findings
    finally:
        workbook.close()


def _is_external_formula(formula: str) -> bool:
    return bool(
        re.search(
            r"https?://|file://|\\\\[^\\]+|\|[^!]+!|"
            r"\b(?:WEBSERVICE|RTD|CALL|EXEC|REGISTER\.ID)\s*\(|"
            r"\[[^\]]+\.(?:xlsx|xlsm|xlsb|xls|csv)\]",
            formula,
            re.IGNORECASE,
        )
    )


def _has_blocking_finding(findings: list[dict[str, object]]) -> bool:
    return any(finding["quality_disposition"] == "BLOCK" for finding in findings)


def _matched_fields(columns: Iterable[object]) -> set[str]:
    normalized = {_normalize_column(column) for column in columns}
    return {
        field
        for field, aliases in FIELD_ALIASES.items()
        if normalized.intersection(aliases)
    }


def _normalize_column(column: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(column).lower())


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _content_fingerprint(frame: pd.DataFrame) -> str:
    payload = {
        "columns": [str(column).strip() for column in frame.columns],
        "rows": [
            [_normalized_cell(value) for value in row]
            for row in frame.reset_index(drop=True).itertuples(index=False, name=None)
        ],
    }
    return hashlib.sha256(
        json.dumps(payload, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _normalized_cell(value: object) -> object:
    if pd.isna(value):
        return None
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return pd.Timestamp(value).isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _finding(code: str, quality_disposition: str, **details: object) -> dict[str, object]:
    finding: dict[str, object] = {
        "code": code,
        "quality_disposition": quality_disposition,
    }
    finding.update(details)
    return finding


def _result(
    *,
    file_path: str,
    file_hash: str | None,
    content_fingerprint: str | None,
    selected_sheet: str | None,
    sheet_candidates: list[dict[str, object]],
    observed_schema: list[str],
    row_count: int,
    findings: list[dict[str, object]],
) -> IntakeResult:
    lifecycle_state = _lifecycle_state(findings)
    quality_disposition = _quality_disposition(findings)
    return IntakeResult(
        file_path=file_path,
        file_hash=file_hash,
        content_fingerprint=content_fingerprint,
        lifecycle_state=lifecycle_state,
        quality_disposition=quality_disposition,
        decision=_decision(lifecycle_state, quality_disposition),
        selected_sheet=selected_sheet,
        sheet_candidates=sheet_candidates,
        observed_schema=observed_schema,
        row_count=row_count,
        reason_codes=[str(finding["code"]) for finding in findings],
        findings=findings,
    )


def _quality_disposition(findings: list[dict[str, object]]) -> str:
    disposition = "PASS"
    for finding in findings:
        candidate = str(finding["quality_disposition"])
        if QUALITY_RANK[candidate] > QUALITY_RANK[disposition]:
            disposition = candidate
    return disposition


def _lifecycle_state(findings: list[dict[str, object]]) -> str:
    codes = {str(finding["code"]) for finding in findings}
    quality_disposition = _quality_disposition(findings)
    if quality_disposition == "BLOCK":
        return "QUARANTINED"
    if codes.intersection(DUPLICATE_CODES):
        return "DUPLICATE"
    if quality_disposition == "REVIEW":
        return "AWAITING_REVIEW"
    return "BRONZE_ACCEPTED"


def _decision(lifecycle_state: str, quality_disposition: str) -> str:
    if lifecycle_state == "DUPLICATE" or quality_disposition == "BLOCK":
        return "BLOCK"
    if quality_disposition == "REVIEW":
        return "REVIEW_REQUIRED"
    if quality_disposition == "WARN":
        return "CONTINUE_WITH_WARNINGS"
    return "CONTINUE"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validate submitted Excel workbooks.")
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--output-dir", default=None, type=Path)
    args = parser.parse_args(argv)

    results = [validate_workbook(path) for path in args.workbooks]
    if args.output_dir:
        for result in results:
            print(write_result(result, args.output_dir))
    else:
        payload = [result.to_dict() for result in results]
        print(json.dumps(payload[0] if len(payload) == 1 else payload, indent=2))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
