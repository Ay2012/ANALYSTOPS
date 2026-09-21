from __future__ import annotations

import json
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from tests.helpers import read_json, write_clean_submission
from analystops.datasets.corruptions import corrupt_submission
from analystops.ingestion.review import (
    ReviewResolutionError,
    main as review_main,
    reassess_workbook,
    write_review_request,
)
from analystops.ingestion.validate import validate_workbook
from analystops.transformations.operations import (
    APPROVED_OPERATIONS,
    allowed_operations,
)


def resolution_for(result, resolutions: list[dict[str, object]]) -> dict[str, object]:
    return {
        "file_hash": result.file_hash,
        "reviewed_by": "analyst@example.com",
        "reviewed_at": "2026-09-17T14:00:00Z",
        "resolutions": resolutions,
    }


class ReviewWorkflowTests(unittest.TestCase):
    def test_approved_operation_registry_defines_review_policy(self) -> None:
        expected = {
            "ambiguous_transaction_sheets": ("select_sheet",),
            "renamed_required_columns": ("map_columns",),
            "quantity_parse_failures": ("confirm_numeric_format",),
            "price_parse_failures": ("confirm_numeric_format",),
            "date_parse_failures": ("confirm_date_format",),
            "non_iso_date_strings": ("confirm_date_format",),
            "exact_duplicate_rows": (
                "confirm_valid_duplicates",
                "deduplicate_in_silver",
            ),
            "empty_transaction_sheet": ("confirm_zero_activity",),
            "unexpectedly_low_row_count": ("confirm_expected_volume",),
            "formula_cells": ("materialize_values_in_silver",),
        }

        self.assertEqual(
            {code: allowed_operations(code) for code in expected},
            expected,
        )
        self.assertEqual(
            APPROVED_OPERATIONS["map_columns"].required_parameters,
            ("mapping",),
        )
        self.assertEqual(APPROVED_OPERATIONS["map_columns"].approval, "AUTOMATIC")
        self.assertTrue(APPROVED_OPERATIONS["map_columns"].executes_in_silver)
        self.assertEqual(APPROVED_OPERATIONS["deduplicate_in_silver"].risk, "HIGH")
        self.assertEqual(
            APPROVED_OPERATIONS["deduplicate_in_silver"].approval,
            "HUMAN",
        )

    def test_review_cli_persists_the_reassessed_result(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            path = root / "low_volume.xlsx"
            resolution_path = root / "resolution.json"
            output_dir = root / "results"
            write_clean_submission(path)
            initial = validate_workbook(path, baseline_row_count=10)
            resolution_path.write_text(
                json.dumps(
                    resolution_for(
                        initial,
                        [
                            {
                                "finding_code": "unexpectedly_low_row_count",
                                "action": "confirm_expected_volume",
                            }
                        ],
                    )
                )
            )

            output = StringIO()
            with redirect_stdout(output):
                exit_code = review_main(
                    [
                        str(path),
                        "--baseline-row-count",
                        "10",
                        "--resolution",
                        str(resolution_path),
                        "--output-dir",
                        str(output_dir),
                    ]
                )
            result_path = Path(output.getvalue().strip())
            persisted = read_json(result_path)

        self.assertEqual(exit_code, 0)
        self.assertEqual(output.getvalue().strip(), str(result_path))
        self.assertEqual(persisted["lifecycle_state"], "BRONZE_ACCEPTED")
        self.assertEqual(persisted["decision"], "CONTINUE_WITH_WARNINGS")

    def test_low_row_count_can_be_specifically_resolved(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "low_volume.xlsx"
            review_dir = Path(tmpdir) / "reviews"
            write_clean_submission(path)
            initial = validate_workbook(path, baseline_row_count=10)
            request_path = write_review_request(initial, review_dir)
            request = read_json(request_path)
            resolution = resolution_for(
                initial,
                [
                    {
                        "finding_code": "unexpectedly_low_row_count",
                        "action": "confirm_expected_volume",
                        "note": "Confirmed for this reporting period.",
                    }
                ],
            )

            result = reassess_workbook(
                path,
                resolution,
                baseline_row_count=10,
            )

        self.assertEqual(request["status"], "AWAITING_REVIEW")
        self.assertEqual(
            request["review_items"][0]["allowed_actions"],
            ["confirm_expected_volume"],
        )
        self.assertEqual(result.lifecycle_state, "BRONZE_ACCEPTED")
        self.assertEqual(result.quality_disposition, "WARN")
        self.assertEqual(result.decision, "CONTINUE_WITH_WARNINGS")
        finding = next(
            item
            for item in result.findings
            if item["code"] == "unexpectedly_low_row_count"
        )
        self.assertEqual(
            finding["review_resolution"]["reviewed_by"],
            "analyst@example.com",
        )

    def test_partial_resolution_remains_awaiting_review(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "two_findings.xlsx"
            write_clean_submission(path)
            workbook = load_workbook(path)
            workbook["Transactions"]["C2"] = "=1+1"
            workbook.save(path)
            workbook.close()
            initial = validate_workbook(path, baseline_row_count=10)
            resolution = resolution_for(
                initial,
                [
                    {
                        "finding_code": "unexpectedly_low_row_count",
                        "action": "confirm_expected_volume",
                    }
                ],
            )

            result = reassess_workbook(
                path,
                resolution,
                baseline_row_count=10,
            )

        self.assertEqual(result.lifecycle_state, "AWAITING_REVIEW")
        findings = {item["code"]: item for item in result.findings}
        self.assertEqual(
            findings["unexpectedly_low_row_count"]["quality_disposition"],
            "WARN",
        )
        self.assertEqual(findings["formula_cells"]["quality_disposition"], "REVIEW")

    def test_renamed_columns_require_an_explicit_mapping(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            clean_path = Path(tmpdir) / "clean.xlsx"
            write_clean_submission(clean_path)
            corrupted = corrupt_submission(
                clean_path,
                scenario="renamed_columns",
                seed=1,
                output_dir=Path(tmpdir) / "corrupted",
            )
            initial = validate_workbook(corrupted.output_path)
            mapping = {
                renamed: canonical
                for canonical, renamed in corrupted.details["renamed_columns"].items()
            }
            resolution = resolution_for(
                initial,
                [
                    {
                        "finding_code": "renamed_required_columns",
                        "action": "map_columns",
                        "details": {"mapping": mapping},
                    }
                ],
            )

            result = reassess_workbook(corrupted.output_path, resolution)

        self.assertEqual(result.lifecycle_state, "BRONZE_ACCEPTED")
        self.assertEqual(result.quality_disposition, "WARN")

    def test_numeric_format_must_match_the_operation_contract(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "currency.xlsx"
            write_clean_submission(path)
            frame = pd.read_excel(path, sheet_name="Transactions")
            frame["Price"] = frame["Price"].map(lambda value: f"${value}")
            frame.to_excel(path, sheet_name="Transactions", index=False)
            initial = validate_workbook(path)
            resolution = resolution_for(
                initial,
                [
                    {
                        "finding_code": "price_parse_failures",
                        "action": "confirm_numeric_format",
                        "details": {"format": "probably-money"},
                    }
                ],
            )

            with self.assertRaisesRegex(ReviewResolutionError, "currency.*number"):
                reassess_workbook(path, resolution)

    def test_human_can_select_between_equally_plausible_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "ambiguous.xlsx"
            source = Path(tmpdir) / "source.xlsx"
            write_clean_submission(source)
            rows = pd.read_excel(source, sheet_name="Transactions")
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                rows.to_excel(writer, sheet_name="Alpha", index=False)
                rows.to_excel(writer, sheet_name="Beta", index=False)
            initial = validate_workbook(path)
            resolution = resolution_for(
                initial,
                [
                    {
                        "finding_code": "ambiguous_transaction_sheets",
                        "action": "select_sheet",
                        "details": {"sheet_name": "Beta"},
                    }
                ],
            )

            result = reassess_workbook(path, resolution)

        self.assertEqual(result.lifecycle_state, "BRONZE_ACCEPTED")
        self.assertEqual(result.selected_sheet, "Beta")

    def test_changed_workbook_invalidates_resolution(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "changed.xlsx"
            write_clean_submission(path)
            initial = validate_workbook(path, baseline_row_count=10)
            resolution = resolution_for(
                initial,
                [
                    {
                        "finding_code": "unexpectedly_low_row_count",
                        "action": "confirm_expected_volume",
                    }
                ],
            )
            workbook = load_workbook(path)
            workbook["Transactions"]["C2"] = "Changed after review"
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ReviewResolutionError, "hash"):
                reassess_workbook(path, resolution, baseline_row_count=10)

    def test_quarantined_workbook_cannot_be_approved_by_review(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "invalid.xlsx"
            write_clean_submission(path)
            frame = pd.read_excel(path, sheet_name="Transactions").drop(columns="Invoice")
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                frame.to_excel(writer, sheet_name="Transactions", index=False)
            initial = validate_workbook(path)
            resolution = resolution_for(initial, [])

            with self.assertRaisesRegex(ReviewResolutionError, "Quarantined"):
                reassess_workbook(path, resolution)


if __name__ == "__main__":
    unittest.main()
